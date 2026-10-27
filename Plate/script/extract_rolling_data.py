"""
ETL Pipeline for Plate Sequential Rolling Data
Extracts necessary features for YS, TS, EL machine learning training.

Features extracted:
  - Identification & Traceability (Batch, Slab, Heat, Roll Date, Spec, Rolling Type, Normalizing Status)
  - 1. Slab Dimension (Grade, Thickness, Width, Length, Weight)
  - 2. Plate Dimension (Thickness, Width, Length, Width Finish, Length Finish, Qty)
  - 3. Chemistry (C, Mn, Si, Al, S, P, Nb, V, Mo, Ti, Cu, Ni, Cr, N, CE)
  - 4. Transfer Bar Thickness (TB Thickness, TB Ratio)
  - 5. Roughing Mill Parameters (Charge Time, Discharge Time, Retention Time, Soaking Temp, Temp Pass 2, RM Finish Temp, Broadsizing Pass, Roughing Pass, Pass Turning, Rolling Strategy)
  - 6. Finishing Mill Parameters (FM Temp Before Rolling, FM Finish Temp, Passes FM, Thickness n-1, Last Pass Thick, Last Pass Drafting)
  - 7. Accelerated Cooling Setup (Laminar Header, Temperature Post Laminar)
  - Target Mechanical Properties (Tensile Before N: YS, TS, EL; Tensile After N: YS, TS, EL)

Uses dynamic Category + Header detection to withstand column location shifts.
Pure Python based (python-calamine engine with openpyxl fallback) - No Excel COM.
"""

import os
import sys
import time
import argparse
import logging
import re
from datetime import date, datetime
import pandas as pd

# Set UTF-8 encoding for standard output
if sys.stdout.encoding.lower() != 'utf-8':
    sys.stdout.reconfigure(encoding='utf-8')

# Ensure script directory is on sys.path
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
if SCRIPT_DIR not in sys.path:
    sys.path.insert(0, SCRIPT_DIR)

from column_schema import SCHEMA, normalize_text

logging.basicConfig(
    level=logging.INFO,
    format='[%(asctime)s] [%(levelname)s] %(message)s',
    datefmt='%Y-%m-%d %H:%M:%S'
)
logger = logging.getLogger("ETL_Rolling")


def find_matching_columns(row0, row2, schema):
    """
    Dynamically maps raw columns to schema definitions by matching:
    1. Category in Row 0 (forward filled)
    2. Header in Row 2
    """
    categories = []
    last_cat = ""
    for val in row0:
        if val is not None and str(val).strip():
            last_cat = normalize_text(val)
        categories.append(last_cat)

    headers = [normalize_text(h) for h in row2]
    total_cols = len(headers)

    mapping = []
    unmatched = []

    for field in schema:
        field_name = field["name"]
        cat_pats = field.get("cat_patterns", [])
        hdr_pats = field.get("header_patterns", [])
        dtype = field.get("dtype", "str")
        group = field.get("group", "General")

        matched_idx = None

        # Priority 1: Match both Category and Header pattern
        for idx in range(total_cols):
            cat = categories[idx]
            hdr = headers[idx]
            cat_ok = any(cp in cat for cp in cat_pats) if cat_pats else True
            hdr_ok = any(re.search(hp, hdr) for hp in hdr_pats)

            if cat_ok and hdr_ok:
                matched_idx = idx
                break

        # Priority 2: Fallback to unique Header pattern match across all columns
        if matched_idx is None:
            candidates = []
            for idx in range(total_cols):
                hdr = headers[idx]
                if any(re.search(hp, hdr) for hp in hdr_pats):
                    candidates.append(idx)
            if len(candidates) == 1:
                matched_idx = candidates[0]

        if matched_idx is not None:
            mapping.append({
                "field_name": field_name,
                "group": group,
                "raw_col_idx": matched_idx,
                "raw_category": row0[matched_idx] or categories[matched_idx],
                "raw_header": row2[matched_idx],
                "dtype": dtype
            })
        else:
            unmatched.append(field_name)

    return mapping, unmatched


def clean_value(val, dtype):
    """Clean and cast individual cell value based on expected dtype."""
    if val is None:
        return None
    val_str = str(val).strip()
    if val_str in ("", "-", "--", "None", "null", "NULL", "#N/A", "#VALUE!", "#REF!", "#DIV/0!"):
        return None

    if dtype == "float":
        try:
            return float(val_str.replace(",", ""))
        except (ValueError, TypeError):
            # Sometimes values have trailing units or comments
            m = re.search(r"[-+]?\d*\.?\d+", val_str)
            if m:
                try:
                    return float(m.group(0))
                except ValueError:
                    return None
            return None

    elif dtype == "int":
        try:
            return int(float(val_str.replace(",", "")))
        except (ValueError, TypeError):
            return None

    elif dtype == "str":
        # Format dates/times if passed as datetime objects
        if isinstance(val, (datetime, date)):
            return val.strftime("%Y-%m-%d")
        return val_str

    return val_str


def load_data_from_excel(filepath, sheet_name="2026"):
    """
    Loads raw sheet rows using fast python-calamine with fallback to openpyxl.
    Resilient to whitespace in sheet names or slight naming differences.
    No COM used.
    """
    logger.info(f"Loading Excel file: {filepath}")
    t0 = time.time()

    # Try python_calamine
    try:
        from python_calamine import CalamineWorkbook
        wb = CalamineWorkbook.from_path(filepath)
        
        target_sheet = None
        for s in wb.sheet_names:
            if s.strip().lower() == sheet_name.strip().lower():
                target_sheet = s
                break
        if target_sheet is None:
            for s in wb.sheet_names:
                if sheet_name.strip().lower() in s.lower():
                    target_sheet = s
                    break
        if target_sheet is None:
            target_sheet = wb.sheet_names[0]
            logger.warning(f"Sheet '{sheet_name}' not found. Defaulting to first sheet: '{target_sheet}'")

        sheet = wb.get_sheet_by_name(target_sheet)
        rows = sheet.to_python()
        logger.info(f"Loaded {len(rows)} rows from sheet [{target_sheet}] using python-calamine in {time.time() - t0:.2f}s")
        return rows
    except Exception as e:
        logger.warning(f"Could not load with python-calamine ({e}), attempting openpyxl fallback...")

    # Fallback openpyxl
    import openpyxl
    wb = openpyxl.load_workbook(filepath, data_only=True, read_only=True)
    target_sheet = None
    for s in wb.sheetnames:
        if s.strip().lower() == sheet_name.strip().lower():
            target_sheet = s
            break
    if target_sheet is None:
        for s in wb.sheetnames:
            if sheet_name.strip().lower() in s.lower():
                target_sheet = s
                break
    if target_sheet is None:
        target_sheet = wb.sheetnames[0]
        logger.warning(f"Sheet '{sheet_name}' not found. Defaulting to first sheet: '{target_sheet}'")

    sheet = wb[target_sheet]
    rows = []
    for r in sheet.iter_rows(values_only=True):
        rows.append(list(r))
    logger.info(f"Loaded {len(rows)} rows from sheet [{target_sheet}] using openpyxl in {time.time() - t0:.2f}s")
    return rows


def process_rolling_data(input_file, sheet_name="2026", output_file=None):
    """Executes the full ETL extraction pipeline."""
    if not os.path.exists(input_file):
        raise FileNotFoundError(f"Input file not found: {input_file}")

    if output_file is None:
        base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
        output_file = os.path.join(base_dir, "csv", f"plate_rolling_{sheet_name}_cleaned.csv")

    os.makedirs(os.path.dirname(output_file), exist_ok=True)

    rows = load_data_from_excel(input_file, sheet_name=sheet_name)
    if len(rows) < 4:
        raise ValueError("Excel sheet has insufficient rows for headers and data.")

    row0 = rows[0]  # Category
    row2 = rows[2]  # Headers
    data_rows = rows[3:]  # Actual records

    logger.info("Detecting dynamic column mappings based on headers and categories...")
    mapping, unmatched = find_matching_columns(row0, row2, SCHEMA)

    logger.info(f"Successfully matched {len(mapping)} out of {len(SCHEMA)} schema columns.")
    if unmatched:
        logger.warning(f"Unmatched schema columns: {unmatched}")

    # Build clean extracted dataset
    col_names = [m["field_name"] for m in mapping]
    indices = [m["raw_col_idx"] for m in mapping]
    dtypes = [m["dtype"] for m in mapping]

    processed_data = []
    valid_record_count = 0

    for r in data_rows:
        # Check if record has any meaningful data (e.g. Batch ID or Slab ID present)
        if not r or all(c is None or str(c).strip() == "" for c in r):
            continue

        row_dict = {}
        for col_name, idx, dt in zip(col_names, indices, dtypes):
            raw_val = r[idx] if idx < len(r) else None
            row_dict[col_name] = clean_value(raw_val, dt)

        # Ensure at least Batch ID or Slab ID exists
        if row_dict.get("Batch_ID") or row_dict.get("Slab_ID"):
            processed_data.append(row_dict)
            valid_record_count += 1

    df = pd.DataFrame(processed_data)

    # Cast integer columns with NaNs safely to pandas nullable Int64
    for m in mapping:
        c_name = m["field_name"]
        if m["dtype"] == "int" and c_name in df.columns:
            df[c_name] = pd.to_numeric(df[c_name], errors='coerce').astype('Int64')

    logger.info(f"Saving {len(df)} cleaned records to CSV: {output_file}")
    df.to_csv(output_file, index=False, encoding='utf-8')

    # Display Summary
    print("\n" + "=" * 80)
    print("PLATE SEQUENTIAL ROLLING RESULT - ETL SUMMARY REPORT")
    print("=" * 80)
    print(f"Source Excel File : {input_file}")
    print(f"Target Sheet      : [{sheet_name}]")
    print(f"Destination CSV   : {output_file}")
    print(f"Total Rows Extracted : {len(df):,}")
    print(f"Columns Extracted    : {len(df.columns)}")
    print("-" * 80)
    print("COLUMN GROUPS BREAKDOWN:")
    groups = {}
    for m in mapping:
        g = m["group"]
        groups.setdefault(g, []).append(m["field_name"])

    for g_name, f_list in groups.items():
        print(f"  * {g_name} ({len(f_list)} columns):")
        for f in f_list:
            non_null_cnt = df[f].notna().sum()
            non_null_pct = non_null_cnt / len(df) * 100 if len(df) > 0 else 0
            print(f"      - {f:30s} : {non_null_cnt:6,d} non-null ({non_null_pct:5.1f}%)")

    print("=" * 80)
    print("[SUCCESS] ETL process completed successfully.")
    print("=" * 80 + "\n")

    return output_file


def find_latest_raw_excel(raw_dir):
    """
    Finds the primary or newest .xlsx in raw directory.
    Ignores temporary Excel lock files (starting with ~$ and ending with .tmp).
    """
    if not os.path.exists(raw_dir):
        return None

    default_file = os.path.join(raw_dir, "PLATE SEQUENTIAL ROLLING RESULT - 2026.xlsx")
    
    excel_files = [
        os.path.join(raw_dir, f) for f in os.listdir(raw_dir)
        if f.lower().endswith((".xlsx", ".xlsm")) and not f.startswith("~$")
    ]
    if not excel_files:
        return None

    # If the default file exists and has the highest or tied timestamp, use it
    excel_files.sort(key=lambda p: os.path.getmtime(p), reverse=True)
    if os.path.exists(default_file) and os.path.getmtime(default_file) >= os.path.getmtime(excel_files[0]):
        return default_file

    return excel_files[0]


def sync_if_needed(repo_root=None, force=False):
    """
    Checks if raw Excel file is newer than cleaned CSV.
    If newer or CSV missing, automatically runs process_rolling_data.
    Returns True if extraction was triggered, False if CSV is already up-to-date.
    """
    if repo_root is None:
        repo_root = os.path.dirname(SCRIPT_DIR)

    raw_dir = os.path.join(repo_root, "raw")
    csv_file = os.path.join(repo_root, "csv", "plate_rolling_2026_cleaned.csv")

    raw_file = find_latest_raw_excel(raw_dir)
    if raw_file is None:
        if os.path.exists(csv_file):
            logger.info("No raw Excel file found in raw/; utilizing existing cleaned CSV.")
            return False
        raise FileNotFoundError(f"No Excel (.xlsx) file found in {raw_dir} and cleaned CSV is missing.")

    excel_mtime = os.path.getmtime(raw_file)
    csv_exists = os.path.exists(csv_file)
    csv_mtime = os.path.getmtime(csv_file) if csv_exists else 0

    if force or (not csv_exists) or (excel_mtime > csv_mtime):
        reason = "force requested" if force else ("CSV missing" if not csv_exists else "raw Excel modified after CSV")
        logger.info(f"[DATA SYNC] Triggering ETL ({reason}): {os.path.basename(raw_file)}")
        process_rolling_data(raw_file, sheet_name="2026", output_file=csv_file)
        return True
    else:
        logger.info(f"[DATA SYNC] CSV is already up-to-date with {os.path.basename(raw_file)} (modified: {datetime.fromtimestamp(excel_mtime).strftime('%Y-%m-%d %H:%M:%S')}).")
        return False


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="ETL for Plate Sequential Rolling Result Excel")
    parser.add_argument("--input", "-i", type=str, default=None,
                        help="Input Excel file path (default: auto-detected from raw/)")
    parser.add_argument("--sheet", "-s", type=str, default="2026",
                        help="Sheet name to extract (default: 2026)")
    parser.add_argument("--output", "-o", type=str, default=None,
                        help="Output CSV path (default: csv/plate_rolling_<sheet>_cleaned.csv)")
    parser.add_argument("--sync", action="store_true",
                        help="Only extract if raw Excel is newer than CSV or CSV is missing")
    parser.add_argument("--force", action="store_true",
                        help="Force re-extraction even if CSV exists")

    args = parser.parse_args()

    repo_root = os.path.dirname(SCRIPT_DIR)
    if args.sync or not args.force:
        if args.input is None and args.output is None:
            sync_if_needed(repo_root=repo_root, force=args.force)
            sys.exit(0)

    # Manual specific path run
    raw_dir = os.path.join(repo_root, "raw")
    input_path = args.input if args.input else find_latest_raw_excel(raw_dir)
    if input_path and not os.path.isabs(input_path):
        input_path = os.path.join(repo_root, input_path)
    output_path = args.output if (args.output is None or os.path.isabs(args.output)) else os.path.join(repo_root, args.output)

    process_rolling_data(input_path, sheet_name=args.sheet, output_file=output_path)
